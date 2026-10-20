# -*- coding: utf-8 -*-
"""把 PS01 传感器的使用说明.docx 全部内容导出成纯文本，便于阅读。"""
import sys
import io
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

path = r"D:\文档\deepseek-harness\default-workspace\ps01_manual.docx"
doc = Document(path)

print("=" * 78)
print("段落内容")
print("=" * 78)
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        style = para.style.name if para.style is not None else "?"
        marker = f"[{style}] " if style not in ("Normal", "Body Text", "正文") else ""
        print(f"{marker}{text}")

print()
print("=" * 78)
print("表格内容")
print("=" * 78)
for ti, table in enumerate(doc.tables):
    print(f"\n--- 表 {ti + 1} （{len(table.rows)} 行 x {len(table.columns)} 列）---")
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
        if any(cells):
            print(" | ".join(cells))

print()
print("=" * 78)
print("统计")
print("=" * 78)
print(f"段落数: {len(doc.paragraphs)}")
print(f"表格数: {len(doc.tables)}")
print(f"内嵌图片数: {len(doc.inline_shapes)}")
