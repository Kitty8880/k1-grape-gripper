# -*- coding: utf-8 -*-
"""
把 Markdown 求助文档转成 Word (.docx)，方便发给厂商/导师。

支持的语法：标题、表格、代码块、无序/有序列表、加粗、水平线、普通段落。
"""
import os
import re
import sys
import io

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

SRC = r"D:\文档\deepseek-harness\default-workspace\grape-gripper\docs\PS01通信问题-求助描述.md"
DST = r"D:\文档\deepseek-harness\default-workspace\PS01通信问题-求助描述.docx"

CN_FONT = "Microsoft YaHei"
MONO_FONT = "Consolas"


def set_run_font(run, name=CN_FONT, size=10.5, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(paragraph, text, base_size=10.5):
    """处理 **加粗** 和 `代码` 内联语法。"""
    # 按 **粗体** 和 `代码` 切分
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            set_run_font(r, size=base_size)
        token = m.group()
        if token.startswith("**"):
            r = paragraph.add_run(token[2:-2])
            set_run_font(r, size=base_size, bold=True)
        else:
            r = paragraph.add_run(token[1:-1])
            set_run_font(r, MONO_FONT, size=base_size - 1)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        set_run_font(r, size=base_size)


def main() -> int:
    if not os.path.isfile(SRC):
        print(f"源文件不存在：{SRC}")
        return 1

    with open(SRC, "r", encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    # 默认样式
    style = doc.styles["Normal"]
    style.font.name = CN_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 空行
        if not stripped:
            i += 1
            continue

        # 水平线
        if re.fullmatch(r"-{3,}", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run("─" * 40)
            set_run_font(r, size=9)
            r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            i += 1
            continue

        # 代码块
        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # 跳过结束 ```
            code_text = "\n".join(code_lines)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            # 背景色（用底纹）
            r = p.add_run(code_text)
            set_run_font(r, MONO_FONT, size=8.5)
            r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            from docx.oxml import OxmlElement
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "F2F2F2")
            p._p.get_or_add_pPr().append(shd)
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            p = doc.add_paragraph()
            size = {1: 16, 2: 14, 3: 12, 4: 11}.get(level, 10.5)
            p.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            set_run_font(r, size=size, bold=True)
            r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            i += 1
            continue

        # 表格
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            # 过滤分隔行（|---|---|）
            rows = []
            for tl in table_lines:
                cells = [c.strip() for c in tl.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    continue
                rows.append(cells)

            if not rows:
                continue

            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for ri, row in enumerate(rows):
                for ci in range(ncols):
                    cell = table.cell(ri, ci)
                    text = row[ci] if ci < len(row) else ""
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    add_inline(p, text, base_size=9)
                    if ri == 0:
                        for r in p.runs:
                            r.bold = True
            doc.add_paragraph()
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, m.group(1))
            i += 1
            continue

        # 有序列表
        m = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, m.group(1))
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            r = p.add_run(stripped.lstrip("> ").strip())
            set_run_font(r, size=10)
            r.italic = True
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        add_inline(p, stripped)
        i += 1

    doc.save(DST)
    size = os.path.getsize(DST)
    print(f"已生成：{DST}")
    print(f"大小：{size} 字节（{size/1024:.1f} KB）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
