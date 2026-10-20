# -*- coding: utf-8 -*-
"""导出 docx 全文（段落 + 表格），用于阅读其他助手的方案。"""
import sys
import io
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

path = r"D:\文档\deepseek-harness\default-workspace\ps01_other.docx"
doc = Document(path)

print("=" * 74)
print("段落")
print("=" * 74)
for p in doc.paragraphs:
    t = p.text.strip()
    if t:
        style = p.style.name if p.style is not None else ""
        prefix = f"[{style}] " if style and style not in ("Normal", "Body Text", "正文") else ""
        print(f"{prefix}{t}")

print()
print("=" * 74)
print("表格")
print("=" * 74)
for ti, table in enumerate(doc.tables):
    print(f"\n--- 表 {ti+1} ({len(table.rows)}行 x {len(table.columns)}列) ---")
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
        if any(cells):
            print(" | ".join(cells))

print()
print(f"统计: 段落 {len(doc.paragraphs)}, 表格 {len(doc.tables)}, 图片 {len(doc.inline_shapes)}")
